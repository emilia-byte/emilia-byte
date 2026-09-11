"""
Unit tests for script.py's account-inventory validators: the regex
patterns and the row-level validate_sheet() logic (blank-row skipping,
optional fields, error messages). No browser/Playwright involved.
"""

import pytest
from openpyxl import Workbook

from script import (
    EMAIL_PATTERN,
    FA_PATTERN,
    NAME_PATTERN,
    PASSWORD_PATTERN,
    PHONE_PATTERN,
    UID_PATTERN,
    SHEET_CONFIGS,
    validate_sheet,
)


# ── Regex patterns ────────────────────────────────────────────────────────

@pytest.mark.parametrize(
    "value,expected",
    [
        ("ABC_ABC_123", True),
        ("ABC_ABC_A123", True),
        ("abc_abc_123", False),   # lowercase not allowed
        ("ABC-ABC-123", False),   # dash instead of underscore
        ("ABCABC123", False),     # missing underscores
        ("AB_CD_12", False),      # too few trailing digits
    ],
)
def test_name_pattern(value, expected):
    assert bool(NAME_PATTERN.match(value)) is expected


@pytest.mark.parametrize(
    "value,expected",
    [
        ("123456789", True),        # 9 digits, lower bound
        ("123456789012345", True),  # 15 digits, upper bound
        ("12345678", False),        # 8 digits, too short
        ("1234567890123456", False),  # 16 digits, too long
        ("12345678a", False),
    ],
)
def test_uid_pattern(value, expected):
    assert bool(UID_PATTERN.match(value)) is expected


@pytest.mark.parametrize(
    "value,expected",
    [
        ("Passw0rd", True),     # 8 chars, has upper + digit
        ("Passw0rd!", True),    # allowed special char
        ("password1", False),   # no uppercase
        ("PASSWORD", False),    # no digit
        ("Pass1", False),       # too short
        ("Passw0rd$", False),   # '$' not in allowed charset
    ],
)
def test_password_pattern(value, expected):
    assert bool(PASSWORD_PATTERN.match(value)) is expected


@pytest.mark.parametrize(
    "value,expected",
    [
        ("A1" * 16, True),   # 32 chars, uppercase letters + digits
        ("A1" * 15, False),  # 30 chars, too short
        ("a1" * 16, False),  # lowercase not allowed
    ],
)
def test_fa_pattern(value, expected):
    assert bool(FA_PATTERN.match(value)) is expected


@pytest.mark.parametrize(
    "value,expected",
    [
        ("user@example.com", True),
        ("not-an-email", False),
        ("user@nodot", False),
    ],
)
def test_email_pattern(value, expected):
    assert bool(EMAIL_PATTERN.match(value)) is expected


@pytest.mark.parametrize(
    "value,expected",
    [
        ("+1234567890", True),
        ("1234567", True),   # 7 chars total, meets the {6,} tail minimum
        ("123456", False),   # 6 chars total, one short
        ("abcdefg", False),
    ],
)
def test_phone_pattern(value, expected):
    assert bool(PHONE_PATTERN.match(value)) is expected


# ── validate_sheet ────────────────────────────────────────────────────────

def _make_sheet(rows):
    """rows: list of tuples (name, uid, password, fa_2, number, cdc, exp)
    written starting at row 3, matching the "BFL Info" column layout
    (B=name, E=uid, F=password, G=fa_2, H=number, I=cdc, J=exp)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "BFL Info"
    for offset, row in enumerate(rows):
        row_idx = 3 + offset
        name, uid, password, fa_2, number, cdc, exp = row
        ws.cell(row=row_idx, column=2, value=name)
        ws.cell(row=row_idx, column=5, value=uid)
        ws.cell(row=row_idx, column=6, value=password)
        ws.cell(row=row_idx, column=7, value=fa_2)
        ws.cell(row=row_idx, column=8, value=number)
        ws.cell(row=row_idx, column=9, value=cdc)
        ws.cell(row=row_idx, column=10, value=exp)
    return ws


VALID_ROW = (
    "ABC_ABC_123",
    "12345678901234",
    "Passw0rd",
    "A1" * 16,
    "+15551234567",
    "yes",
    "2026",
)


def test_validate_sheet_accepts_a_fully_valid_row():
    ws = _make_sheet([VALID_ROW])
    errors = validate_sheet(ws, SHEET_CONFIGS["BFL Info"])
    assert errors == []


def test_validate_sheet_skips_fully_blank_rows():
    ws = _make_sheet([VALID_ROW, (None, None, None, None, None, None, None), VALID_ROW])
    errors = validate_sheet(ws, SHEET_CONFIGS["BFL Info"])
    assert errors == []


def test_validate_sheet_reports_invalid_uid_with_row_and_column():
    bad_row = ("ABC_ABC_123", "not-a-uid", "Passw0rd", "A1" * 16, "+15551234567", "yes", "2026")
    ws = _make_sheet([bad_row])
    errors = validate_sheet(ws, SHEET_CONFIGS["BFL Info"])

    assert len(errors) == 1
    assert "Row 3" in errors[0]
    assert "UID" in errors[0]
    assert "col E" in errors[0]


def test_validate_sheet_reports_missing_non_optional_field():
    bad_row = ("ABC_ABC_123", "12345678901234", "Passw0rd", "A1" * 16, "+15551234567", None, "2026")
    ws = _make_sheet([bad_row])
    errors = validate_sheet(ws, SHEET_CONFIGS["BFL Info"])

    assert len(errors) == 1
    assert "Missing CDC" in errors[0]


def test_validate_sheet_optional_password_skipped_when_blank():
    row = ("ABC_ABC_123", "123456789012345", None, "A1" * 16, "user@example.com", "prov", None)
    wb = Workbook()
    ws = wb.active
    ws.title = "New VProfiles for BFL"
    ws.cell(row=2, column=3, value=row[0])
    ws.cell(row=2, column=4, value=row[1])
    ws.cell(row=2, column=5, value=row[2])  # password: blank, optional field
    ws.cell(row=2, column=6, value=row[3])
    ws.cell(row=2, column=7, value=row[4])
    ws.cell(row=2, column=9, value=row[5])

    errors = validate_sheet(ws, SHEET_CONFIGS["New VProfiles for BFL"])
    assert errors == []
